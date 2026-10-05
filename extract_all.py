import openpyxl, sys, io, json, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
SRC = r"C:\Users\rizky\Downloads\Copy of Wedding Planner for Reina.xlsx"
OUT = r"C:\Users\rizky\Documents\Dev\wedding-planner\seed"
import os
os.makedirs(OUT, exist_ok=True)
wb = openpyxl.load_workbook(SRC, data_only=True)

def c(v):
    return "" if v is None else str(v).strip()

def num(x):
    try: return float(x)
    except: return 0.0

# ============ SESERAHAN ============
ws = wb['List Seserahan']
seserahan = []
section = "LAIN-LAIN"
SECTIONS = {"ALAT SHOLAT","PAKAIAN","ACCESSORIES","HAIR CARE","BODY CARE & TOILETRIES","SKINCARE","MAKEUP"}
for r in ws.iter_rows(min_row=3, max_row=200, values_only=True):
    v = [c(x) for x in r]
    no, detail, brand, qty, link, budget, checklist, actual = (v + [""]*8)[:8]
    if detail.upper() in SECTIONS and not no:
        section = detail.upper(); continue
    if detail.upper() == "DETAIL":
        continue
    if not detail or detail.lower().startswith("total"):
        continue
    seserahan.append({
        "category": section,
        "detail": detail,
        "brand": brand,
        "qty": int(num(qty)) or 1,
        "link": link,
        "budget": num(budget),
        "checklist": checklist.lower() == "true",
        "actually_harga": num(actual),
    })
with open(os.path.join(OUT,"seserahan.json"),"w",encoding="utf-8") as f:
    json.dump(seserahan,f,ensure_ascii=False,indent=2)
print("seserahan:", len(seserahan), "sections:", sorted({s['category'] for s in seserahan}))

# ============ ENGAGEMENT BUDGET ============
def parse_budget(sheet):
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=1, max_row=200, values_only=True))
    items=[]; section="LAIN-LAIN"
    for r in rows:
        v=[c(x) for x in r]
        detail,status,vendor,total,ket,t1,t2,t3,t4,*_=(v+[""]*15)[:9]
        if detail=="DETAIL": continue
        if detail and not status and not vendor and not total:  # section heading
            if detail.upper()==detail and detail not in ("","DETAIL"):
                section=detail; continue
        if not detail: continue
        items.append({
            "category":section,"detail":detail,"status":status or "Not Yet",
            "vendor":vendor,"total_biaya":num(total),"keterangan":ket,
            "termin1":num(t1),"termin2":num(t2),"termin3":num(t3),"termin4":num(t4),
        })
    return items

eng = parse_budget('Engagement Budget')
wed = parse_budget('Wedding Budget')
with open(os.path.join(OUT,"engagement_budget.json"),"w",encoding="utf-8") as f:
    json.dump(eng,f,ensure_ascii=False,indent=2)
with open(os.path.join(OUT,"wedding_budget.json"),"w",encoding="utf-8") as f:
    json.dump(wed,f,ensure_ascii=False,indent=2)
print("engagement items:", len(eng), "wedding items:", len(wed))

# ============ ADMINISTRASI & DOKUMEN ============
ws = wb['Administrasi & Dokumen']
admin=[]; section="LAIN-LAIN"
for r in ws.iter_rows(min_row=3, max_row=200, values_only=True):
    v=[c(x) for x in r]
    sec, name, progress, cpp, cpw, details = (v+[""]*6)[:6]
    # heading rows: col A has letter section and col B is section title
    if sec and name and not progress:
        section=f"{sec} - {name}"; continue
    if not name and not sec: continue
    if name.lower() in ("daftar persiapan","progress","details"): continue
    if not name: continue
    is_doc = cpp or cpw
    admin.append({
        "section":section,
        "item": name,
        "progress": progress if not is_doc else "",
        "cpp_complete": cpp.lower()=="true",
        "cpw_complete": cpw.lower()=="true",
        "is_document": bool(is_doc),
        "details": details,
    })
with open(os.path.join(OUT,"administrations.json"),"w",encoding="utf-8") as f:
    json.dump(admin,f,ensure_ascii=False,indent=2)
print("admin:", len(admin))

# ============ VENDOR CHECKLIST ============
ws = wb['Vendor Checklist']
vc=[]; section="LAIN-LAIN"
for r in ws.iter_rows(min_row=4, max_row=200, values_only=True):
    v=[c(x) for x in r]
    activity, deadline, countdown, checklist, ket = (v+[""]*5)[:5]
    if activity in ("Booking Vendor","Things To Do","Administrasi Pernikahan","Rangkaian Acara"):
        section=activity; continue
    if not activity or activity.startswith("dst") or activity.startswith("Kegiatan"): continue
    def d(x):
        x=str(x)[:10]
        return x if re.match(r"\d{4}-\d{2}-\d{2}",x) else ""
    vc.append({
        "category":section,"activity_name":activity,"deadline_date":d(deadline),
        "is_completed":checklist.lower()=="true","notes":ket,
    })
with open(os.path.join(OUT,"vendor_checklists.json"),"w",encoding="utf-8") as f:
    json.dump(vc,f,ensure_ascii=False,indent=2)
print("vendor checklist:", len(vc), sorted({x['category'] for x in vc}))

# ============ WEDDING CASHFLOW ============
ws = wb['Wedding Cashflow']
cf=[]
for r in ws.iter_rows(min_row=4, max_row=40, values_only=True):
    v=[c(x) for x in r]
    no, tgl, ket, inc_cpp, inc_cpw, exp_cpp, exp_cpw = (v+[""]*7)[:7]
    if not ket or ket in ("Tabungan",): 
        tgl_ok = re.match(r"\d{4}-\d{2}-\d{2}", tgl)
        if not (ket and tgl_ok): 
            if ket=="Tabungan": pass
            else: continue
    date = str(tgl)[:10] if re.match(r"\d{4}-\d{2}-\d{2}", str(tgl)) else ""
    cf.append({"no": int(num(no)) if num(no) else len(cf)+1, "date":date, "description":ket,
               "income_cpp":num(inc_cpp),"income_cpw":num(inc_cpw),
               "expense_cpp":num(exp_cpp),"expense_cpw":num(exp_cpw)})
cf=[x for x in cf if x["description"]]
with open(os.path.join(OUT,"cashflow.json"),"w",encoding="utf-8") as f:
    json.dump(cf,f,ensure_ascii=False,indent=2)
print("cashflow:", len(cf))
print(json.dumps(cf,ensure_ascii=False))
