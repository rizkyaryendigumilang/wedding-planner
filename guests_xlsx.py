import openpyxl, sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

path = r"C:\Users\rizky\Downloads\Copy of Wedding Planner for Reina.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)

# ---- Guest sheets: find header row then pull only rows with a real name
for sh in ['Daftar Tamu Kevin', 'Daftar Tamu Aurene', 'Daftar Tamu Kevin & Aurene']:
    ws = wb[sh]
    print("\n#### SHEET:", sh, ws.max_row, ws.max_column)
    header = None
    count = 0
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        vals = ["" if c is None else str(c).strip() for c in r]
        joined = " | ".join(v for v in vals if v)
        if not joined:
            continue
        # detect header
        if header is None and any('Nama Tamu' in v for v in vals):
            header = vals
            print("HEADER:", header)
            continue
        if header and any('Nama Tamu' in v for v in vals):
            print("(sheet header again)", joined)
            continue
        # guest row: col idx1 (B) is name
        name = vals[1] if len(vals) > 1 else ""
        if name and not name.replace('.', '').replace(',', '').isdigit():
            count += 1
            print("G:", joined)
    print("REAL GUEST ROWS:", count)
