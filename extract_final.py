import openpyxl, sys, io, json, re, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
SRC = r"C:\Users\rizky\Downloads\Copy of Wedding Planner for Reina.xlsx"
OUT = r"C:\Users\rizky\Documents\Dev\wedding-planner\seed"
os.makedirs(OUT, exist_ok=True)
wb = openpyxl.load_workbook(SRC, data_only=True)

def c(v): return "" if v is None else str(v).strip()
def num(x):
    try: return float(x)
    except: return 0.0
def d10(x):
    s = str(x)[:10]
    return s if re.match(r"^\d{4}-\d{2}-\d{2}$", s) else ""

# ================= SESERAHAN =================
ws = wb['List Seserahan']
out=[]; section="LAIN-LAIN"
SECTIONS={"ALAT SHOLAT","PAKAIAN","ACCESSORIES","HAIR CARE","BODY CARE & TOILETRIES","SKINCARE","MAKEUP"}
for r in range(6, 120):
    A,B,C,D,E,F,G,H = (c(ws.cell(r,col).value) for col in range(1,9))
    if A.upper() in SECTIONS and not B:
        section=A.upper(); continue
    if not B or B.lower().startswith(("detail","total")): continue
    out.append({"category":section,"detail":B,"brand":C,"qty":int(num(D)) or 1,
                "link":E,"budget":num(F),"checklist":G.lower()=="true","actually_harga":num(H)})
json.dump(out, open(os.path.join(OUT,"seserahan.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print("seserahan:", len(out), "| sections:", {s['category'] for s in out})

# ================= ENGAGEMENT / WEDDING BUDGET =================
def budgets(sheet):
    ws=wb[sheet]; out=[]; section="LAIN-LAIN"
    for r in range(6, 120):
        A,B,C,D,E,F,G,H,I = (c(ws.cell(r,col).value) for col in range(1,10))
        if A and not B and not C and not D:      # section heading
            section=A; continue
        if not A or A.lower().startswith(("detail","total")): continue
        out.append({"category":section,"detail":A,"status":B or "Not Yet","vendor":C,
                    "total_biaya":num(D),"keterangan":E,"termin1":num(F),
                    "termin2":num(G),"termin3":num(H),"termin4":num(I)})
    return out
eng=budgets('Engagement Budget'); wed=budgets('Wedding Budget')
json.dump(eng, open(os.path.join(OUT,"engagement_budget.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
json.dump(wed, open(os.path.join(OUT,"wedding_budget.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print("engagement:", len(eng), "| wedding:", len(wed))
print("  engagement sections:", sorted({x['category'] for x in eng}))

# ================= ADMINISTRASI =================
ws=wb['Administrasi & Dokumen']; out=[]; section="LAIN-LAIN"
for r in range(6, 120):
    A,B,C,D,E,F = (c(ws.cell(r,col).value) for col in range(1,7))
    # section heading: A = letter, B = title, C empty
    if A and B and not C and len(A)<=2:
        section=f"{A}. {B}"; continue
    if not B or B.lower() in ("daftar persiapan","progress","details"): continue
    is_doc = bool(D or E) and not C   # CPP/CPW columns present
    out.append({"section":section,"item":B,
                "progress": C if not is_doc else "",
                "cpp_complete": D.lower()=="true","cpw_complete": E.lower()=="true",
                "is_document": is_doc,"details":F})
json.dump(out, open(os.path.join(OUT,"administrations.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print("admin:", len(out), "| sections:", sorted({x['section'] for x in out}))

# ================= VENDOR CHECKLIST =================
ws=wb['Vendor Checklist']; out=[]; section="LAIN-LAIN"
CATS={"Booking Vendor","Things To Do","Administrasi Pernikahan","Rangkaian Acara"}
for r in range(4, 140):
    A,B,C,D,E = (c(ws.cell(r,col).value) for col in range(1,6))
    if A in CATS: section=A; continue
    if not A or A.lower().startswith(("kegiatan","dst")): continue
    out.append({"category":section,"activity_name":A,"deadline_date":d10(B) or d10(C),
                "is_completed":(D or C).lower()=="true" if (D or C) else False,"notes":E})
json.dump(out, open(os.path.join(OUT,"vendor_checklists.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print("vendor checklist:", len(out))
for cat in CATS: print("  ", cat, sum(1 for x in out if x['category']==cat))

# ================= CASHFLOW =================
ws=wb['Wedding Cashflow']; out=[]
for r in range(4, 40):
    A,B,C,D,E,F,G = (c(ws.cell(r,col).value) for col in range(1,8))
    if not C or C.lower() in ("keterangan",): continue
    if not re.match(r"^\d+", A): continue
    out.append({"date":d10(B),"description":C,"income_cpp":num(D),"income_cpw":num(E),
                "expense_cpp":num(F),"expense_cpw":num(G)})
json.dump(out, open(os.path.join(OUT,"cashflow.json"),"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print("cashflow:", len(out))
