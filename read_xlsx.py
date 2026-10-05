import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

path = r"C:\Users\rizky\Downloads\Copy of Wedding Planner for Reina.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("\n" + "=" * 90)
    print(f"SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("=" * 90)
    for row in ws.iter_rows(values_only=True):
        # skip fully empty rows
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        cells = [("" if c is None else str(c).replace("\n", " / ").strip()) for c in row]
        # trim trailing empties
        while cells and cells[-1] == "":
            cells.pop()
        print(" | ".join(cells))
